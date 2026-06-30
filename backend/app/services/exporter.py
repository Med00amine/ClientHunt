"""CSV export. Files are never overwritten, formula injection is neutralised, UTF-8 with BOM for Excel."""
from __future__ import annotations

import collections
import csv
import hashlib
import os
from datetime import datetime
from pathlib import Path

from sqlalchemy import select
from sqlalchemy.orm import Session

from ..config import get_settings
from ..models import (BusinessContact, Company, Evidence, ExportRecord, Opportunity, ProposalDraft,
                      ResearchRun, RunOpportunity, utcnow)
from ..textutil import csv_safe
from . import leads

OPP_COLUMNS = [
    "run_id", "discovered_at", "source_platform", "source_type", "opportunity_id", "title", "company_name",
    "company_website", "original_url", "posted_at", "deadline", "engagement_type", "remote_status",
    "location_requirements", "budget_min", "budget_max", "currency", "rate_type", "description",
    "required_skills", "matched_skills", "portfolio_evidence", "fit_score", "fit_explanation",
    "missing_requirements", "company_research_summary", "research_sources", "contact_name", "contact_role",
    "business_email", "contact_source_url", "contact_verified_status", "outreach_eligibility",
    "recommended_action", "proposal_status", "application_status", "last_verified_at",
    "hiring_ease_score", "hiring_signals", "apply_method",
]
EMAIL_COLUMNS = [
    "run_id", "company_name", "company_website", "business_email", "email_type", "contact_purpose",
    "outreach_eligibility", "eligibility_basis", "email_source_url", "email_found_at", "email_context",
    "verified_status", "listing_title", "listing_url", "source_platform", "apply_method", "hiring_ease_score",
    "hiring_signals", "fit_score", "engagement_type", "location_requirements", "posted_at", "opportunity_id",
]
CONTACT_COLUMNS = ["contact_id", "name", "role", "company_name", "business_email", "is_generic_address",
                   "source_type", "source_url", "retrieved_at", "verified_status", "outreach_eligibility",
                   "eligibility_basis"]


def _dt(v: datetime | None) -> str:
    return v.strftime("%Y-%m-%dT%H:%M:%SZ") if v else ""


def _pick_contact(db: Session, opp: Opportunity) -> BusinessContact | None:
    """Only contacts with recorded provenance and not suppressed: the one tied to this listing first
    (careers/application addresses before general ones), else one for the same company."""
    base = select(BusinessContact).where(BusinessContact.source_url != "", BusinessContact.suppressed.is_(False),
                                         BusinessContact.outreach_eligibility != "not_eligible")
    rank = (BusinessContact.outreach_eligibility != "eligible", BusinessContact.id)
    c = db.scalar(base.where(BusinessContact.opportunity_id == opp.id).order_by(*rank).limit(1))
    if c is None and opp.company_id:
        c = db.scalar(base.where(BusinessContact.company_id == opp.company_id).order_by(*rank).limit(1))
    return c


def opportunity_rows(db: Session, run_id: int) -> list[dict]:
    links = db.scalars(select(RunOpportunity).where(RunOpportunity.run_id == run_id)
                       .order_by(RunOpportunity.id)).all()
    rows, seen = [], set()
    for link in links:
        if link.opportunity_id in seen:
            continue
        seen.add(link.opportunity_id)
        o = db.get(Opportunity, link.opportunity_id)
        company = db.get(Company, o.company_id) if o.company_id else None
        facts = []
        if company:
            facts = db.scalars(select(Evidence).where(Evidence.company_id == company.id,
                                                      Evidence.kind == "fact")).all()
        c = _pick_contact(db, o)
        elig = "unknown"
        if c:
            elig = c.outreach_eligibility if c.outreach_eligibility in ("eligible", "not_eligible") else "review_required"
        draft = db.scalar(select(ProposalDraft).where(ProposalDraft.opportunity_id == o.id)
                          .order_by(ProposalDraft.id.desc()).limit(1))
        score = o.score_override if o.score_override is not None else o.fit_score
        rows.append({
            "run_id": run_id, "discovered_at": _dt(link.discovered_at), "source_platform": o.source_key,
            "source_type": o.source_type, "opportunity_id": o.id, "title": o.title,
            "company_name": o.company_name, "company_website": company.website if company else "",
            "original_url": o.original_url, "posted_at": _dt(o.posted_at), "deadline": _dt(o.deadline),
            "engagement_type": o.engagement_type, "remote_status": o.remote_status,
            "location_requirements": o.location_requirements,
            "budget_min": "" if o.budget_min is None else o.budget_min,
            "budget_max": "" if o.budget_max is None else o.budget_max,
            "currency": o.currency, "rate_type": o.rate_type, "description": o.description,
            "required_skills": "; ".join(o.required_skills or []),
            "matched_skills": "; ".join(o.matched_skills or []),
            "portfolio_evidence": "; ".join(e["title"] for e in (o.portfolio_evidence or []) if e.get("verified")),
            "fit_score": score, "fit_explanation": o.fit_explanation,
            "missing_requirements": "; ".join(o.missing_requirements or []),
            "company_research_summary": " | ".join(f.claim for f in facts[:5]),
            "research_sources": "; ".join(sorted({f.source_url for f in facts if f.source_url})),
            "contact_name": c.name if c else "", "contact_role": c.role if c else "",
            "business_email": c.email if c else "", "contact_source_url": c.source_url if c else "",
            "contact_verified_status": c.verified_status if c else "", "outreach_eligibility": elig,
            "recommended_action": o.recommended_action, "proposal_status": o.proposal_status or ("draft" if draft else "none"),
            "application_status": o.application_status, "last_verified_at": _dt(o.last_verified_at),
            "hiring_ease_score": o.ease_score, "hiring_signals": "; ".join(o.ease_signals or []),
            "apply_method": o.apply_method,
        })
    return rows


def _write_unique(directory: Path, prefix: str, columns: list[str], rows: list[dict]) -> tuple[Path, str]:
    directory.mkdir(parents=True, exist_ok=True)
    stamp = utcnow().strftime("%Y-%m-%d_%H-%M-%S")
    n = 0
    while True:
        name = f"{prefix}_{stamp}{'' if n == 0 else f'_{n + 1}'}.csv"
        path = directory / name
        try:
            fd = os.open(path, os.O_WRONLY | os.O_CREAT | os.O_EXCL, 0o644)  # never overwrite
            break
        except FileExistsError:
            n += 1
    try:
        with os.fdopen(fd, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore", quoting=csv.QUOTE_MINIMAL)
            w.writeheader()
            for r in rows:
                w.writerow({k: csv_safe(r.get(k, "")) for k in columns})
        digest = hashlib.sha256(path.read_bytes()).hexdigest()
    except Exception:
        path.unlink(missing_ok=True)  # no partial file left behind
        raise
    return path, digest


def export_run(db: Session, run: ResearchRun, commit: bool = True) -> ExportRecord:
    """Create the opportunity CSV for a run and record it. On failure a failed record is stored."""
    s = get_settings()
    try:
        rows = opportunity_rows(db, run.id)
        path, digest = _write_unique(s.exports_dir, "opportunities", OPP_COLUMNS, rows)
        rec = ExportRecord(run_id=run.id, kind="opportunities", filename=path.name, path=str(path),
                           row_count=len(rows), sha256=digest, status="ok", run_status=run.status)
        run.export_status = "ok"
    except Exception as e:  # noqa: BLE001
        rec = ExportRecord(run_id=run.id, kind="opportunities", status="failed", error=str(e)[:500],
                           run_status=run.status)
        run.export_status = "failed"
    db.add(rec)
    if commit:
        db.commit()
    return rec


def export_contacts(db: Session) -> ExportRecord:
    """Optional contacts CSV: only eligible, non-suppressed contacts that have recorded provenance."""
    s = get_settings()
    try:
        contacts = db.scalars(select(BusinessContact).where(
            BusinessContact.outreach_eligibility == "eligible", BusinessContact.suppressed.is_(False),
            BusinessContact.source_url != "")).all()
        rows = [{"contact_id": c.id, "name": c.name, "role": c.role, "company_name": c.company_name,
                 "business_email": c.email, "is_generic_address": c.is_generic, "source_type": c.source_type,
                 "source_url": c.source_url, "retrieved_at": _dt(c.retrieved_at),
                 "verified_status": c.verified_status, "outreach_eligibility": c.outreach_eligibility,
                 "eligibility_basis": c.eligibility_basis} for c in contacts]
        path, digest = _write_unique(s.exports_dir, "contacts", CONTACT_COLUMNS, rows)
        rec = ExportRecord(kind="contacts", filename=path.name, path=str(path), row_count=len(rows),
                           sha256=digest, status="ok")
    except Exception as e:  # noqa: BLE001
        rec = ExportRecord(kind="contacts", status="failed", error=str(e)[:500])
    db.add(rec)
    db.commit()
    return rec


def _email_rows(db: Session, contacts: list[BusinessContact], run_id: int | str = "") -> list[dict]:
    rows = []
    for c in contacts:
        o = db.get(Opportunity, c.opportunity_id) if c.opportunity_id else None
        company = db.get(Company, c.company_id) if c.company_id else (db.get(Company, o.company_id) if o and o.company_id else None)
        rows.append({
            "run_id": run_id, "company_name": c.company_name or (o.company_name if o else ""),
            "company_website": company.website if company else "", "business_email": c.email,
            "email_type": "role / generic mailbox" if c.is_generic else "individual address (as published)",
            "contact_purpose": c.contact_purpose, "outreach_eligibility": c.outreach_eligibility,
            "eligibility_basis": c.eligibility_basis, "email_source_url": c.source_url,
            "email_found_at": _dt(c.retrieved_at), "email_context": c.source_note,
            "verified_status": c.verified_status, "listing_title": o.title if o else "",
            "listing_url": o.original_url if o else "", "source_platform": o.source_key if o else "",
            "apply_method": o.apply_method if o else "", "hiring_ease_score": o.ease_score if o else "",
            "hiring_signals": "; ".join(o.ease_signals or []) if o else "",
            "fit_score": (o.score_override if o.score_override is not None else o.fit_score) if o else "",
            "engagement_type": o.engagement_type if o else "", "location_requirements": o.location_requirements if o else "",
            "posted_at": _dt(o.posted_at) if o else "", "opportunity_id": o.id if o else "",
        })
    rows.sort(key=lambda r: (r["outreach_eligibility"] != "eligible", -(r["hiring_ease_score"] or 0 if r["hiring_ease_score"] != "" else 0)))
    return rows


def _usable_contacts(stmt):
    return stmt.where(BusinessContact.source_url != "", BusinessContact.suppressed.is_(False),
                      BusinessContact.outreach_eligibility != "not_eligible")


def _prospect_rows(db: Session, contacts: list[BusinessContact]) -> list[dict]:
    """Companies the user chose to approach. Mailboxes are role addresses with the page they were found on."""
    from .prospects import OFFERS, SOURCE
    from ..models import Application
    by_opp: dict[int, list[BusinessContact]] = {}
    for c in contacts:
        if c.opportunity_id:
            by_opp.setdefault(c.opportunity_id, []).append(c)
    rows = []
    for o in db.scalars(select(Opportunity).where(Opportunity.source_key == SOURCE, Opportunity.review_status != "excluded")
                        .order_by(Opportunity.ease_score.desc(), Opportunity.id)).all():
        meta = o.meta or {}
        cs = sorted(by_opp.get(o.id, []), key=lambda c: (c.outreach_eligibility != "eligible", c.id))
        draft = db.scalar(select(ProposalDraft).where(ProposalDraft.opportunity_id == o.id).order_by(ProposalDraft.id.desc()).limit(1))
        app = db.scalar(select(Application).where(Application.opportunity_id == o.id))
        why = meta.get("evidence_quote", "")
        if meta.get("evidence_url"):
            why = f"{why} ({meta['evidence_url']})".strip()
        rows.append({
            "company": o.title, "website": o.original_url, "category": meta.get("category", "company"),
            "offer": OFFERS.get(meta.get("offer", ""), {}).get("label", ""), "why_added": why,
            "public_signals": "; ".join(o.ease_signals or []), "signal_score": o.ease_score,
            "business_mailboxes": "; ".join(c.email for c in cs), "mailbox_eligibility": ", ".join(sorted({c.outreach_eligibility for c in cs})),
            "mailbox_source_urls": "; ".join(sorted({c.source_url for c in cs})), "research_status": o.research_status,
            "draft_status": draft.status if draft else "none", "contact_status": app.status if app else "not contacted",
            "last_verified_at": o.last_verified_at, "recommended_action": o.recommended_action})
    return rows


def run_emails(db: Session, run_id: int) -> list[BusinessContact]:
    opp_ids = select(RunOpportunity.opportunity_id).where(RunOpportunity.run_id == run_id)
    return list(db.scalars(_usable_contacts(select(BusinessContact).where(BusinessContact.opportunity_id.in_(opp_ids)))
                           .order_by(BusinessContact.id)).all())


def export_run_emails(db: Session, run: ResearchRun, commit: bool = True) -> ExportRecord:
    """Emails CSV for a run: only addresses the company published (with source URL and date)."""
    s = get_settings()
    try:
        rows = _email_rows(db, run_emails(db, run.id), run.id)
        path, digest = _write_unique(s.exports_dir, "emails", EMAIL_COLUMNS, rows)
        rec = ExportRecord(run_id=run.id, kind="emails", filename=path.name, path=str(path), row_count=len(rows),
                           sha256=digest, status="ok", run_status=run.status)
    except Exception as e:  # noqa: BLE001
        rec = ExportRecord(run_id=run.id, kind="emails", status="failed", error=str(e)[:500], run_status=run.status)
        run.export_status = "failed"
    db.add(rec)
    if commit:
        db.commit()
    return rec


def export_run_bundle(db: Session, run: ResearchRun) -> tuple[ExportRecord, ExportRecord]:
    """Opportunities CSV + emails CSV, committed together with the run's final status."""
    a = export_run(db, run, commit=False)
    b = export_run_emails(db, run, commit=False)
    db.commit()
    return a, b


def export_all_emails(db: Session) -> ExportRecord:
    """Every stored address that has a source URL, is not suppressed and not marked not_eligible."""
    s = get_settings()
    try:
        contacts = db.scalars(_usable_contacts(select(BusinessContact)).order_by(BusinessContact.id)).all()
        rows = _email_rows(db, list(contacts))
        path, digest = _write_unique(s.exports_dir, "emails_all", EMAIL_COLUMNS, rows)
        rec = ExportRecord(kind="emails", filename=path.name, path=str(path), row_count=len(rows), sha256=digest, status="ok")
    except Exception as e:  # noqa: BLE001
        rec = ExportRecord(kind="emails", status="failed", error=str(e)[:500])
    db.add(rec)
    db.commit()
    return rec


# ====================================================================== ONE FILE with everything (all runs)
LEAD_MIN_FIT = 35.0   # the Leads sheet only lists opportunities that fit you at least this well
COMBINED_OPP_COLUMNS = [
    "opportunity_id", "source_platform", "source_type", "title", "company_name", "company_website", "original_url",
    "posted_at", "deadline", "engagement_type", "remote_status", "location_requirements", "location_check", "budget_min", "budget_max",
    "currency", "rate_type", "fit_score", "hiring_ease_score", "apply_method", "business_emails", "email_eligibility",
    "fit_explanation", "hiring_signals", "required_skills", "matched_skills", "missing_requirements",
    "portfolio_evidence", "company_research_summary", "research_sources", "recommended_action", "review_status",
    "proposal_status", "application_status", "runs_seen", "first_seen_at", "last_seen_at", "last_verified_at",
    "description",
]
LEAD_COLUMNS = [
    "rank", "company_name", "title", "business_email", "email_eligibility", "email_type", "email_found_at",
    "email_source_url", "email_context", "engagement_type", "location_requirements", "location_check", "budget", "posted_at", "fit_score",
    "hiring_ease_score", "why_it_looks_easy", "listing_url", "source_platform", "recommended_action", "runs_seen",
]
PROSPECT_COLUMNS = ["company", "website", "category", "offer", "why_added", "public_signals", "signal_score", "business_mailboxes",
                    "mailbox_eligibility", "mailbox_source_urls", "research_status", "draft_status", "contact_status",
                    "last_verified_at", "recommended_action"]
RUN_COLUMNS = ["run_id", "search", "status", "started_at", "finished_at", "duration_s", "sources", "fetched", "stored",
               "new", "rejected_irrelevant", "rejected_duplicate", "selected_for_review", "companies_researched",
               "drafts_generated", "emails_found", "errors"]
URL_COLUMNS = {"company_website", "original_url", "email_source_url", "listing_url", "contact_source_url", "website", "mailbox_source_urls"}
DATE_COLUMNS = {"posted_at", "deadline", "first_seen_at", "last_seen_at", "last_verified_at", "email_found_at",
                "started_at", "finished_at"}
_XL_MAX = 32000
_ILLEGAL = __import__("re").compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _budget_text(o: Opportunity) -> str:
    if o.budget_min is None and o.budget_max is None:
        return ""
    lo, hi = o.budget_min, o.budget_max
    rng = f"{lo:g}-{hi:g}" if lo is not None and hi is not None and lo != hi else f"{(hi if hi is not None else lo):g}"
    return f"{rng} {o.currency}{' / ' + o.rate_type if o.rate_type else ''}".strip()


def build_combined(db: Session, include_excluded: bool = False) -> dict:  # noqa: C901
    """Everything from ALL runs as plain Python tables. Opportunities are de-duplicated across runs (one row each,
    with the runs that saw them). Suppressed / not-eligible addresses and addresses without a source are left out.
    The Leads sheet lists only opportunities that fit you (score >= LEAD_MIN_FIT) and are not restricted to other
    countries; the Opportunities and Emails sheets keep everything (see the location_check column)."""
    from ..scoring import effective_geo

    runs = db.scalars(select(ResearchRun).order_by(ResearchRun.id)).all()
    seen_rows = db.execute(select(RunOpportunity.opportunity_id, RunOpportunity.run_id, RunOpportunity.discovered_at)
                           .order_by(RunOpportunity.run_id)).all()
    seen: dict[int, list] = {}
    for oid, rid, at in seen_rows:
        seen.setdefault(oid, []).append((rid, at))
    stmt = select(Opportunity).where(Opportunity.id.in_(list(seen))) if seen else select(Opportunity).where(False)
    opps = [o for o in db.scalars(stmt).all() if include_excluded or o.review_status != "excluded"]
    omitted = len(seen) - len(opps)

    contacts = list(db.scalars(_usable_contacts(select(BusinessContact)).order_by(BusinessContact.id)).all())
    by_opp: dict[int, list[BusinessContact]] = {}
    by_company: dict[int, list[BusinessContact]] = {}
    for c in contacts:
        if c.opportunity_id:
            by_opp.setdefault(c.opportunity_id, []).append(c)
        if c.company_id and c.contact_purpose not in leads.PITCH_PURPOSES:   # a pitch mailbox is not a job-application address
            by_company.setdefault(c.company_id, []).append(c)

    def contacts_for(o: Opportunity) -> list[BusinessContact]:
        got = {c.id: c for c in by_opp.get(o.id, [])}
        for c in by_company.get(o.company_id or -1, []):
            got.setdefault(c.id, c)
        return sorted(got.values(), key=lambda c: (c.outreach_eligibility != "eligible", c.id))

    opp_rows, lead_rows, lead_opp_ids = [], [], set()
    for o in opps:
        company = db.get(Company, o.company_id) if o.company_id else None
        facts = db.scalars(select(Evidence).where(Evidence.company_id == company.id, Evidence.kind == "fact")).all() if company else []
        cs = contacts_for(o)
        runs_seen = sorted({rid for rid, _ in seen.get(o.id, [])})
        times = [at for _, at in seen.get(o.id, [])]
        fit = o.score_override if o.score_override is not None else o.fit_score
        loc_check = effective_geo(o.location_requirements, o.description)
        elig = ", ".join(sorted({c.outreach_eligibility for c in cs}))
        opp_rows.append({
            "opportunity_id": o.id, "source_platform": o.source_key, "source_type": o.source_type, "title": o.title,
            "company_name": o.company_name, "company_website": company.website if company else "",
            "original_url": o.original_url, "posted_at": o.posted_at, "deadline": o.deadline,
            "engagement_type": o.engagement_type, "remote_status": o.remote_status,
            "location_requirements": o.location_requirements, "location_check": loc_check,
            "budget_min": o.budget_min, "budget_max": o.budget_max, "currency": o.currency, "rate_type": o.rate_type, "fit_score": fit, "hiring_ease_score": o.ease_score,
            "apply_method": o.apply_method, "business_emails": "; ".join(c.email for c in cs), "email_eligibility": elig,
            "fit_explanation": o.fit_explanation, "hiring_signals": "; ".join(o.ease_signals or []),
            "required_skills": "; ".join(o.required_skills or []), "matched_skills": "; ".join(o.matched_skills or []),
            "missing_requirements": "; ".join(o.missing_requirements or []),
            "portfolio_evidence": "; ".join(e["title"] for e in (o.portfolio_evidence or []) if e.get("verified")),
            "company_research_summary": " | ".join(f.claim for f in facts[:5]),
            "research_sources": "; ".join(sorted({f.source_url for f in facts if f.source_url})),
            "recommended_action": o.recommended_action, "review_status": o.review_status,
            "proposal_status": o.proposal_status, "application_status": o.application_status,
            "runs_seen": ", ".join(str(r) for r in runs_seen), "first_seen_at": min(times) if times else o.first_seen_at,
            "last_seen_at": max(times) if times else o.last_seen_at, "last_verified_at": o.last_verified_at,
            "description": o.description,
        })
        is_lead = bool(cs) and (fit or 0) >= LEAD_MIN_FIT and loc_check != "restricted" and o.listing_status != "closed"
        if is_lead:
            lead_opp_ids.add(o.id)
        for c in (cs if is_lead else []):
            lead_rows.append({
                "company_name": c.company_name or o.company_name, "title": o.title, "business_email": c.email,
                "email_eligibility": c.outreach_eligibility,
                "email_type": "role / generic mailbox" if c.is_generic else "individual address (as published)",
                "email_found_at": c.retrieved_at, "email_source_url": c.source_url, "email_context": c.source_note,
                "engagement_type": o.engagement_type, "location_requirements": o.location_requirements,
                "location_check": loc_check, "budget": _budget_text(o), "posted_at": o.posted_at, "fit_score": fit, "hiring_ease_score": o.ease_score,
                "why_it_looks_easy": "; ".join(o.ease_signals or []), "listing_url": o.original_url,
                "source_platform": o.source_key, "recommended_action": o.recommended_action,
                "runs_seen": ", ".join(str(r) for r in runs_seen), "_sort": ((fit or 0) + (o.ease_score or 0)) / 2,
            })
    loc_order = {"worldwide": 0, "compatible": 0, "unknown": 1, "restricted": 2}   # open to Tunisia first, restricted last
    opp_rows.sort(key=lambda r: (loc_order.get(r["location_check"], 1), -((r["fit_score"] or 0) + (r["hiring_ease_score"] or 0)) / 2))
    lead_rows.sort(key=lambda r: (r["email_eligibility"] != "eligible", -r["_sort"]))
    for i, r in enumerate(lead_rows, 1):
        r["rank"] = i
        r.pop("_sort", None)

    email_rows = _email_rows(db, contacts)
    prospect_rows = _prospect_rows(db, contacts)
    run_rows = []
    for r in runs:
        emails_in_run = len(run_emails(db, r.id)) if r.status in ("completed", "partial") else 0
        run_rows.append({
            "run_id": r.id, "search": r.profile_name, "status": r.status, "started_at": r.started_at,
            "finished_at": r.finished_at, "duration_s": int((r.finished_at - r.started_at).total_seconds()) if r.started_at and r.finished_at else "",
            "sources": ", ".join(r.sources_queried or []), "fetched": r.discovered, "stored": r.stored, "new": r.new_count,
            "rejected_irrelevant": r.rejected_irrelevant, "rejected_duplicate": r.rejected_duplicate,
            "selected_for_review": r.selected_for_review, "companies_researched": r.companies_researched,
            "drafts_generated": r.drafts_generated, "emails_found": emails_in_run, "errors": " | ".join(r.errors or [])})
    summary = [
        ("Generated (UTC)", utcnow()),
        ("Runs included", len(runs)),
        ("Opportunities (each listed once, even if several runs found it)", len(opp_rows)),
        ("Leads (opportunities with a published email that fit you and are open to Tunisia)", len(lead_opp_ids)),
        ("Emails (company-published, each with its source link)", len(email_rows)),
        ("Prospects (companies you chose to approach; see the Prospects sheet)", len(prospect_rows)),
        ("Excluded opportunities left out", omitted),
        ("By location check", ", ".join(f"{k}: {v}" for k, v in sorted(collections.Counter(r["location_check"] for r in opp_rows).items(),
                                                                         key=lambda kv: loc_order.get(kv[0], 1))) or "none"),
        ("", ""),
        ("How to read it", f"Leads = one row per email for opportunities with fit >= {LEAD_MIN_FIT:g} that are not restricted to other countries, "
                           "best fit + easiest to reach first. Opportunities = everything found, one row each (filter the location_check column). "
                           "Emails = every address with where and when it was found. Runs = the history of your searches."),
        ("location_check", "worldwide / compatible (Europe, Africa, EMEA) / restricted (limited to other countries) / unknown (not stated)."),
        ("eligible", "The company itself published this address to receive applications. Use it only to apply or ask about that role."),
        ("review_required", "General mailbox or free-mail address. Check the rules for the recipient's country before writing."),
        ("Fit / Easy to reach", "Heuristics from the listing text (0-100). They tell you what to read first; they do not predict hiring."),
        ("Sources", "Listings link back to their source. Jobicy, Remotive, Remote OK and Himalayas ask for credit; keep the links."),
        ("Not included", "Freelancer.com projects: its API terms forbid storing its data. Use Find work → Live projects."),
    ]
    return {"summary": summary, "leads": lead_rows, "opportunities": opp_rows, "prospects": prospect_rows, "emails": email_rows,
            "runs": run_rows,
            "counts": {"runs": len(runs), "opportunities": len(opp_rows), "leads": len(lead_opp_ids), "emails": len(email_rows),
                       "prospects": len(prospect_rows)}}


def _xl_clean(v):
    if isinstance(v, str):
        return _ILLEGAL.sub("", v)[:_XL_MAX]
    return v


def build_xlsx(data: dict) -> bytes:
    """One Excel workbook. Text is always stored as text (a value starting with = is never a formula) and only
    http(s) links become clickable."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill, head_font = PatternFill("solid", fgColor="2457D6"), Font(bold=True, color="FFFFFF")
    wide = {"title": 52, "description": 90, "fit_explanation": 70, "company_name": 26, "business_email": 34,
            "business_emails": 38, "email_context": 60, "email_source_url": 40, "original_url": 40, "listing_url": 40,
            "why_it_looks_easy": 60, "hiring_signals": 60, "recommended_action": 46, "location_requirements": 30,
            "company_research_summary": 60, "search": 34, "errors": 50, "eligibility_basis": 50}

    def put(ws, r, c, v, col):
        cell = ws.cell(row=r, column=c)
        v = _xl_clean(v)
        if v is None or v == "":
            return
        cell.value = v
        if isinstance(v, str):
            cell.data_type = "s"            # never a formula, even if it starts with = + - @
            if col in URL_COLUMNS and v.lower().startswith(("http://", "https://")):
                cell.hyperlink = v
                cell.font = Font(color="2457D6", underline="single")
        elif hasattr(v, "year"):
            cell.number_format = "yyyy-mm-dd hh:mm"

    def table(ws, columns, rows):
        for c, name in enumerate(columns, 1):
            h = ws.cell(row=1, column=c, value=name)
            h.fill, h.font = head_fill, head_font
            ws.column_dimensions[get_column_letter(c)].width = wide.get(name, 16)
        for r, row in enumerate(rows, 2):
            for c, name in enumerate(columns, 1):
                put(ws, r, c, row.get(name), name)
        ws.freeze_panes = "A2"
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"

    ws = wb.active
    ws.title = "Summary"
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 58, 110
    ws.cell(row=1, column=1, value="ClientHunt: everything from all runs in one file").font = Font(bold=True, size=14)
    for i, (k, v) in enumerate(data["summary"], 3):
        a, b = ws.cell(row=i, column=1), ws.cell(row=i, column=2)
        a.value, a.data_type = _xl_clean(k), "s"
        if hasattr(v, "year"):
            b.value, b.number_format = v, "yyyy-mm-dd hh:mm"
        else:
            b.value = _xl_clean(v)
            if isinstance(v, str):
                b.data_type = "s"
        a.alignment = b.alignment = Alignment(wrap_text=True, vertical="top")
        if k in ("eligible", "review_required", "How to read it", "location_check", "Fit / Easy to reach", "Sources", "Not included"):
            a.font = Font(bold=True)
    for name, key, cols in (("Leads", "leads", LEAD_COLUMNS), ("Opportunities", "opportunities", COMBINED_OPP_COLUMNS),
                            ("Prospects", "prospects", PROSPECT_COLUMNS), ("Emails", "emails", EMAIL_COLUMNS),
                            ("Runs", "runs", RUN_COLUMNS)):
        table(wb.create_sheet(name), cols, data[key])
    import io
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_combined_csv(data: dict) -> bytes:
    """The same information as ONE flat CSV (one row per opportunity, emails joined with '; ')."""
    import io
    buf = io.StringIO(newline="")
    w = csv.DictWriter(buf, fieldnames=COMBINED_OPP_COLUMNS, extrasaction="ignore", quoting=csv.QUOTE_MINIMAL)
    w.writeheader()
    for r in data["opportunities"]:
        out = {}
        for k in COMBINED_OPP_COLUMNS:
            v = r.get(k, "")
            out[k] = _dt(v) if hasattr(v, "year") else ("" if v is None else v)
        w.writerow({k: csv_safe(v) for k, v in out.items()})
    return ("﻿" + buf.getvalue()).encode("utf-8")
