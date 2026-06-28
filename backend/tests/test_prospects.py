import io

import httpx
import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app import api_prospects
from app.config import get_settings
from app.models import Application, BusinessContact, Company, Evidence, Opportunity, ProposalDraft
from app.services import drafting, exporter, prospects as P
from app.services.safefetch import FetchBlocked, SafeFetcher

HOME = ('<html><head><title>Northwind Docs – paperwork automation</title><meta name="description" content="We automate invoice processing for finance teams with AI."></head>'
        '<body><h1>Northwind Docs</h1><p>Our platform uses machine learning to read invoices and contracts. We also offer workflow automation and an API.</p>'
        '<p>We run on AWS and offer staff augmentation and dedicated teams for agencies.</p></body></html>')
CAREERS = '<html><body><h1>Careers</h1><p>We are hiring engineers. Apply at careers@northwind.example</p></body></html>'
CONTACT = ('<html><body><h1>Contact</h1><p>General: hello@northwind.example</p><p>Press: press@northwind.example</p>'
           '<p>We are always looking for freelance engineers: write to partners@northwind.example to get in touch.</p>'
           '<p>Personal: jane.doe@northwind.example</p></body></html>')


def site(pages=None, status=200):
    pages = {"/": HOME, "/careers": CAREERS, "/contact": CONTACT, **(pages or {})}
    def h(req):
        if req.url.path == "/robots.txt":
            return httpx.Response(404)
        body = pages.get(req.url.path)
        if status != 200:
            return httpx.Response(status)
        return httpx.Response(200, text=body, headers={"content-type": "text/html"}) if body else httpx.Response(404)
    return SafeFetcher(httpx.Client(transport=httpx.MockTransport(h)), resolver=lambda host: ["93.184.216.34"])


# ---------------------------------------------------------------- address validation
@pytest.mark.parametrize("raw,expected", [
    ("https://www.northwind.example/about?x=1", "https://www.northwind.example"), ("northwind.example", "https://northwind.example"),
    ("HTTP://Acme.Example/", "http://acme.example"), ("  https://a.example  ", "https://a.example"),
])
def test_normalize_site_accepts_public_sites_and_keeps_only_the_origin(raw, expected):
    assert P.normalize_site(raw) == expected


@pytest.mark.parametrize("bad", ["", "javascript:alert(1)", "file:///etc/passwd", "http://localhost:8000", "http://127.0.0.1", "http://169.254.169.254/latest",
                                 "http://[::1]/", "https://user:pw@a.example", "https://a b.example", "ftp://x.example", "intranet", "http://10.0.0.5", "http://x.local",
                                 "<script>.example"])
def test_normalize_site_rejects_unsafe_or_private_addresses(bad):
    with pytest.raises(ValueError):
        P.normalize_site(bad)


# ---------------------------------------------------------------- adding
def test_add_prospect_creates_an_opportunity_like_record_and_dedupes_by_site(db):
    a = P.add_prospect(db, "https://www.northwind.example/team", "Northwind Docs", "company")
    b = P.add_prospect(db, "northwind.example", "", "agency", "subcontract_capacity", "https://evidence.example/p", "Offers staff augmentation")
    assert a.id == b.id and db.query(Opportunity).filter_by(source_key="prospect").count() == 1
    assert (a.source_type, a.original_url, a.title) == ("manual", "https://www.northwind.example", "Northwind Docs")
    assert a.meta["category"] == "agency" and a.meta["offer"] == "subcontract_capacity" and a.meta["evidence_url"] == "https://evidence.example/p"
    assert db.query(Company).count() == 1 and db.scalar(select(Company)).website == "https://www.northwind.example"


def test_a_prospect_never_adopts_a_same_named_company_that_has_another_website(db):
    from app.textutil import norm_text
    other = Company(name="Acme", name_norm=norm_text("Acme"), website="https://acme-old.example")
    bare = Company(name="Beta", name_norm=norm_text("Beta"), website="")                      # known only from a job listing
    db.add_all([other, bare]); db.commit()
    a = P.add_prospect(db, "https://acme-ai.example", "Acme", "company")
    b = P.add_prospect(db, "https://beta.example", "Beta", "company")
    assert a.company_id != other.id and db.get(Company, a.company_id).website == "https://acme-ai.example"   # research reads THEIR site
    assert b.company_id == bare.id and db.get(Company, b.company_id).website == "https://beta.example"      # a bare company gets the address
    assert db.get(Company, other.id).website == "https://acme-old.example"


@pytest.mark.parametrize("kw", [dict(category="spam"), dict(offer="nope"), dict(evidence_url="javascript:alert(1)")])
def test_add_prospect_validates_its_inputs(db, kw):
    with pytest.raises(ValueError):
        P.add_prospect(db, "https://a.example", **kw)


# ---------------------------------------------------------------- research
def test_research_keeps_verbatim_evidence_and_classifies_mailboxes_for_a_pitch(db):
    o = P.add_prospect(db, "https://northwind.example", "Northwind Docs", "company")
    res = P.research_prospect(db, o, site())
    db.refresh(o)
    assert res["status"] == "done" and o.research_status == "done"
    facts = db.scalars(select(Evidence).where(Evidence.kind == "fact")).all()
    assert facts and all(not f.approved and f.source_url.startswith("https://northwind.example") for f in facts)
    boxes = {c.email: (c.outreach_eligibility, c.contact_purpose) for c in db.scalars(select(BusinessContact))}
    assert boxes["hello@northwind.example"] == ("review_required", "b2b_mailbox")                # unsolicited pitch: never "eligible"
    assert boxes["careers@northwind.example"] == ("review_required", "b2b_mailbox")
    assert boxes["partners@northwind.example"] == ("eligible", "partner_invited")                # the site invites freelancers there
    assert "press@northwind.example" not in boxes and "jane.doe@northwind.example" not in boxes   # no press address, no named person
    assert all(c.source_url.startswith("https://northwind.example/") and c.retrieved_at for c in db.scalars(select(BusinessContact)))
    assert o.apply_method == "email" and "mailbox" in o.recommended_action.lower()
    assert o.meta["pages"] and o.meta["researched_at"]


def test_signals_are_public_hints_with_points_not_proof(db):
    o = P.add_prospect(db, "https://northwind.example", "Northwind Docs", "company")
    P.research_prospect(db, o, site())
    db.refresh(o)
    s = " | ".join(o.ease_signals)
    for needle in ("AI / machine learning (+20)", "automation, integrations or APIs (+15)", "cloud / AWS (+10)", "delivery capacity",
                   "careers/jobs page", "invites freelancers", "role mailbox"):
        assert needle in s, needle
    assert 0 < o.ease_score <= 100 and o.ease_score == min(100, 20 + 15 + 10 + 15 + 15 + 25 + 10)


def test_offer_is_suggested_from_the_site_unless_the_user_chose_one(db):
    auto = P.add_prospect(db, "https://northwind.example", "Northwind Docs", "company")
    P.research_prospect(db, auto, site())
    db.refresh(auto)
    assert auto.meta["offer"] in ("document_extraction", "workflow_automation") and auto.meta["offer_auto"] is True
    chosen = P.add_prospect(db, "https://chosen.example", "Chosen", "company", "aws_deployment")
    P.research_prospect(db, chosen, site())
    db.refresh(chosen)
    assert chosen.meta["offer"] == "aws_deployment" and chosen.meta["offer_auto"] is False
    agency = P.add_prospect(db, "https://agency.example", "Agency", "agency")
    P.research_prospect(db, agency, site())
    db.refresh(agency)
    assert agency.meta["offer"] == "subcontract_capacity"


def test_network_prospects_use_application_rules_so_a_careers_mailbox_is_eligible(db):
    o = P.add_prospect(db, "https://northwind.example", "Some Network", "network")
    P.research_prospect(db, o, site())
    boxes = {c.email: c.outreach_eligibility for c in db.scalars(select(BusinessContact))}
    assert boxes["careers@northwind.example"] == "eligible" and boxes["hello@northwind.example"] == "review_required"


def test_a_network_is_not_scored_as_if_it_needed_work_done(db):
    o = P.add_prospect(db, "https://northwind.example", "Some Network", "network")
    P.research_prospect(db, o, site())
    db.refresh(o)
    assert o.ease_score == 0 and list(o.ease_signals) == [] and o.research_status == "done"
    company = P.add_prospect(db, "https://other.example", "Other", "company")
    P.research_prospect(db, company, site())
    db.refresh(company)
    assert company.ease_score > 0                                       # companies keep their signs of need


def test_a_network_page_that_says_send_your_resume_here_makes_that_mailbox_eligible(db):
    page = '<html><body><h1>Join</h1><p>To join Some Network, please send your resumé (PDF or DOC only) to info@network.example.</p></body></html>'
    o = P.add_prospect(db, "https://network.example", "Some Network", "network")
    P.research_prospect(db, o, site({"/": page, "/careers": None, "/contact": None}))
    box = db.scalar(select(BusinessContact).where(BusinessContact.email == "info@network.example"))
    assert (box.outreach_eligibility, box.contact_purpose) == ("eligible", "application_invited") and box.source_url == "https://network.example"
    pitch = P.add_prospect(db, "https://agency.example", "Agency", "agency")
    P.research_prospect(db, pitch, site({"/": page.replace("network.example", "agency.example"), "/careers": None, "/contact": None}))
    b2 = db.scalar(select(BusinessContact).where(BusinessContact.email == "info@agency.example"))
    assert (b2.outreach_eligibility, b2.contact_purpose) == ("review_required", "b2b_mailbox")     # to an agency it is only a pitch


def test_unreachable_site_is_recorded_not_raised(db):
    o = P.add_prospect(db, "https://down.example", "Down Co", "company")
    P.research_prospect(db, o, site(status=500))
    db.refresh(o)
    assert o.research_status == "failed" and "Could not read the website" in o.recommended_action
    assert db.query(BusinessContact).count() == 0


def test_research_never_fetches_private_addresses(db):
    called = []
    f = SafeFetcher(httpx.Client(transport=httpx.MockTransport(lambda r: called.append(str(r.url)) or httpx.Response(200, text="x"))))
    f.resolver = lambda host: (_ for _ in ()).throw(FetchBlocked("private"))
    o = P.add_prospect(db, "https://evil.example", "Evil", "company")
    P.research_prospect(db, o, f)
    db.refresh(o)
    assert o.research_status == "failed" and called == []                      # the guard stopped it before any request


@pytest.mark.parametrize("text,invited", [
    ("We are always looking for freelance engineers.", True),
    ("Join our network of freelancers and work on interesting projects.", True),
    ("Freelancers and contractors can contact us about projects.", True),
    ("We welcome subcontractors who want to deliver with us.", True),
    ("Freelancers: weak governance, inconsistent availability, knowledge loss.", False),
    ("We are a long-term partner, not a temporary fix. Contact us.", False),
    ("Looking for a freelancer? Hire our dedicated team instead.", False),
    ("We are not looking for freelancers at the moment.", False),
    ("We never work with freelancers.", False),
    ("Become a partner in our reseller program.", False),
    ("Hire vetted freelancers in 48 hours.", False),
    ("", False),
])
def test_only_a_clear_invitation_counts_as_inviting_freelancers(text, invited):
    from app.services import leads
    assert leads.invites_freelancers(text) is invited


def test_the_mailbox_name_cannot_invite_on_its_own_behalf(db):
    from app.services import leads
    o = P.add_prospect(db, "https://acme.example", "Acme", "company")
    cs = leads.attach_emails(db, o, "Partners: partners@acme.example. Freelance: freelance@acme.example, contractors@acme.example",
                             "website", "https://acme.example/contact", "acme.example", mode="pitch")
    assert {c.outreach_eligibility for c in cs} == {"review_required"}                      # names alone are not invitations
    o2 = P.add_prospect(db, "https://beta.example", "Beta", "company")
    inv = leads.attach_emails(db, o2, "We are hiring contractors:\nwrite to talent@beta.example\n", "website", "https://beta.example/x", "beta.example", mode="pitch")
    assert [c.outreach_eligibility for c in inv] == ["review_required"]                     # the invitation is on another line: review, not eligible


def test_a_site_that_criticises_freelancers_gets_no_freelancer_signal(db):
    page = ('<html><body><p>Freelancers: weak governance, inconsistent availability, knowledge loss. We are a long-term partner, not a temporary fix. '
            'Contact: hello@lsk.example</p></body></html>')
    o = P.add_prospect(db, "https://lsk.example", "LSK", "agency")
    P.research_prospect(db, o, site({"/": page, "/careers": None, "/contact": None}))
    db.refresh(o)
    assert not any("invites freelancers" in s for s in o.ease_signals)


def test_research_many_records_failures_and_reset_stuck_recovers(db):
    from app.db import SessionLocal
    o = P.add_prospect(db, "https://boom.example", "Boom", "company")
    class Boom:
        def __init__(self): raise RuntimeError("no network")
    P.research_many(SessionLocal, [o.id, 99999], Boom)
    db.refresh(o)
    assert o.research_status == "failed" and "RuntimeError" in o.recommended_action
    o.research_status = "researching"; db.commit()
    assert P.reset_stuck(db) == 1
    db.refresh(o)
    assert o.research_status == "not_started"


# ---------------------------------------------------------------- pitch drafts
def approved_fact(db, o):
    f = db.scalar(select(Evidence).where(Evidence.kind == "fact", Evidence.claim.like("Company describes itself%")))
    f.approved = True; db.commit()
    return f


def test_pitch_uses_only_approved_facts_verified_work_and_keeps_the_opt_out(db):
    o = P.add_prospect(db, "https://northwind.example", "Northwind Docs", "company", "document_extraction")
    P.research_prospect(db, o, site())
    d = drafting.generate_draft(db, o)
    assert "Replace the placeholders" in " ".join(d.warnings) and any("No approved company fact" in w for w in d.warnings)
    assert "I read on your website" not in d.body                                    # nothing unapproved is quoted
    assert "if this is not relevant for you, please tell me and i will not write again" in d.body.lower()
    assert "[Your name]" in d.body and "Client Acquisition Studio" in d.body         # only the verified item is cited
    assert "Churn" not in d.body and "Vocality" not in d.body and "RAG project" not in d.body
    f = approved_fact(db, o)
    d2 = drafting.generate_draft(db, o)
    assert f"I read on your website: “{f.excerpt[:180]}”." in d2.body and not any("No approved company fact" in w for w in d2.warnings)
    assert d2.version == 2 and d2.generator == "template" and d2.est_cost_usd == 0 and o.proposal_status == "draft"
    assert any("unsolicited" in w for w in d2.warnings)


def test_pitch_never_makes_claims_of_past_results(db):
    for key in P.OFFERS:
        o = P.add_prospect(db, f"https://{key.replace('_', '-')}.example", key, "network" if key == "network_application" else "company", key)
        body = drafting.generate_draft(db, o).body
        assert drafting.check_claims(body, o, {"https://github.com/Med00amine/neurogrid"}, []) == [] or all(
            "years" not in w and "numeric" not in w and "certification" not in w and "compliment" not in w
            for w in drafting.check_claims(body, o, {"https://github.com/Med00amine/neurogrid"}, [])), key


def test_network_application_has_no_opt_out_sentence_but_asks_for_availability(db):
    o = P.add_prospect(db, "https://network.example", "Some Network", "network")
    body = drafting.generate_draft(db, o).body
    assert "not write again" not in body and "[hours] hours per week" in body and "join your network" in body


def test_pitch_ignores_the_llm_even_when_enabled(db, monkeypatch):
    monkeypatch.setattr(get_settings(), "llm_enabled", True)
    monkeypatch.setattr(get_settings(), "anthropic_api_key", "k")
    o = P.add_prospect(db, "https://northwind.example", "Northwind Docs", "company")
    assert drafting.generate_draft(db, o).generator == "template"                   # a pitch is never machine-generated


# ---------------------------------------------------------------- API
@pytest.fixture
def fake_site(monkeypatch):
    monkeypatch.setattr(api_prospects, "fetcher_factory", lambda: site())


def test_api_adds_researches_lists_and_excludes_prospects_from_job_views(client, db, fake_site):
    r = client.post("/api/prospects", json={"items": [{"website": "northwind.example", "name": "Northwind Docs", "category": "company",
                                                      "evidence_url": "https://x.example/a", "evidence_quote": "Offers staff augmentation"},
                                                     {"website": "javascript:alert(1)"}, {"website": "http://localhost"}]})
    assert r.status_code == 200 and len(r.json()["created"]) == 1 and len(r.json()["errors"]) == 2
    assert {e["website"] for e in r.json()["errors"]} == {"javascript:alert(1)", "http://localhost"}
    lst = client.get("/api/prospects").json()
    p = lst["items"][0]
    assert lst["total"] == 1 and p["research_status"] == "done" and p["need_score"] > 0                  # background research ran
    assert {c["email"] for c in p["contacts"]} >= {"hello@northwind.example", "partners@northwind.example"}
    assert p["evidence_quote"] == "Offers staff augmentation" and p["offer_label"]
    assert client.get("/api/prospects?category=agency").json()["total"] == 0
    assert client.get("/api/prospects?only_with_mailbox=true").json()["total"] == 1
    # kept out of job lists, leads and counts
    assert client.get("/api/opportunities").json()["total"] == 0
    assert client.get("/api/leads?has_email=false&hide_restricted=false").json()["items"] == []
    ov = client.get("/api/overview").json()
    assert ov["opportunities_total"] == 0 and ov["emails_found"] == 0 and ov["apply_by_email_leads"] == 0 and ov["prospects"] == 1


def test_api_meta_patch_research_again_and_delete(client, db, fake_site):
    pid = client.post("/api/prospects", json={"items": [{"website": "https://northwind.example", "name": "NW"}], "research": False}).json()["created"][0]["id"]
    assert client.get("/api/prospects").json()["items"][0]["research_status"] == "not_started"
    m = client.get("/api/prospects/meta").json()
    assert set(m["offers"]) == set(P.OFFERS) and set(m["categories"]) == {"company", "agency", "network"}
    assert client.patch(f"/api/prospects/{pid}", json={"offer": "nope"}).status_code == 422
    assert client.patch(f"/api/prospects/{pid}", json={"offer": "aws_deployment"}).json()["offer"] == "aws_deployment"
    assert client.patch(f"/api/prospects/{pid}", json={"category": "agency"}).json()["category"] == "agency"
    assert client.post(f"/api/prospects/{pid}/research").json()["research_status"] in ("researching", "done")
    assert client.get("/api/prospects").json()["items"][0]["research_status"] == "done"
    # drafts, approval and "contacted" use the normal endpoints
    d = client.post(f"/api/opportunities/{pid}/drafts").json()
    assert d["warnings"] and client.post(f"/api/drafts/{d['id']}/approve", json={"acknowledge_warnings": True}).status_code == 200
    assert client.post(f"/api/opportunities/{pid}/applications", json={}).status_code == 200
    assert client.delete(f"/api/prospects/{pid}").status_code == 409                       # an application exists
    assert client.get("/api/prospects").json()["items"][0]["application_status"] == "submitted"


def test_the_generic_research_endpoint_reads_a_prospect_in_pitch_mode(client, db, monkeypatch):
    from app.services import safefetch
    monkeypatch.setattr("app.api.SafeFetcher", lambda: site())
    pid = client.post("/api/prospects", json={"items": [{"website": "https://northwind.example"}], "research": False}).json()["created"][0]["id"]
    assert client.post(f"/api/opportunities/{pid}/research").status_code == 200
    boxes = {c.email: c.contact_purpose for c in db.scalars(select(BusinessContact))}
    assert boxes["hello@northwind.example"] == "b2b_mailbox" and boxes["careers@northwind.example"] == "b2b_mailbox"   # not 'careers_mailbox'


def test_api_delete_removes_the_prospect_its_mailboxes_and_drafts(client, db, fake_site):
    pid = client.post("/api/prospects", json={"items": [{"website": "https://northwind.example"}]}).json()["created"][0]["id"]
    client.post(f"/api/opportunities/{pid}/drafts")
    assert client.delete(f"/api/prospects/{pid}").status_code == 200
    db.expire_all()
    assert db.query(Opportunity).filter_by(source_key="prospect").count() == 0
    assert db.query(BusinessContact).count() == 0 and db.query(ProposalDraft).count() == 0
    assert db.query(Company).count() == 1                                                  # the company record stays
    assert client.delete(f"/api/prospects/{pid}").status_code == 404


def test_api_validation_limits_and_auth(client, monkeypatch):
    assert client.post("/api/prospects", json={"items": []}).status_code == 422
    assert client.post("/api/prospects", json={"items": [{"website": "a.example"}] * 26}).status_code == 422
    assert client.post("/api/prospects", json={"items": [{"website": "a.example", "category": "spam"}]}).status_code == 422
    assert client.get("/api/prospects?category=bad").status_code == 422
    assert client.delete("/api/prospects/99999").status_code == 404
    monkeypatch.setattr(get_settings(), "app_password", "s3cret")
    for method, url in (("get", "/api/prospects"), ("post", "/api/prospects"), ("get", "/api/prospects/meta"), ("delete", "/api/prospects/1")):
        assert getattr(client, method)(url).status_code == 401, url


# ---------------------------------------------------------------- export and maintenance
def test_prospects_get_their_own_sheet_and_stay_out_of_leads_and_opportunities(client, db, fake_site):
    from .helpers import FakeAdapter, raw
    from .test_csv_and_runs import go
    go(db, {"fake": FakeAdapter([raw(1, description="Python FastAPI LLM AWS Docker. To apply, email your resume to jobs@acme.example " * 2, company_name="Acme")])})
    client.post("/api/prospects", json={"items": [{"website": "https://northwind.example", "name": "Northwind Docs", "evidence_quote": "Sells staff augmentation",
                                                  "evidence_url": "https://x.example/p"}]})
    wb = load_workbook(io.BytesIO(exporter.build_xlsx(exporter.build_combined(db))))
    head = [c.value for c in wb["Prospects"][1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in wb["Prospects"].iter_rows(min_row=2)]
    assert len(rows) == 1 and rows[0]["company"] == "Northwind Docs" and rows[0]["website"] == "https://northwind.example"
    assert "Sells staff augmentation (https://x.example/p)" == rows[0]["why_added"]
    assert "partners@northwind.example" in rows[0]["business_mailboxes"] and rows[0]["contact_status"] == "not contacted"
    assert rows[0]["mailbox_source_urls"].startswith("https://northwind.example")
    opps = [r[3].value for r in wb["Opportunities"].iter_rows(min_row=2)]
    assert "Northwind Docs" not in opps and len(opps) == 1                                  # only the real listing
    assert {r[3].value for r in wb["Leads"].iter_rows(min_row=2)} == {"jobs@acme.example"}   # the prospect's mailboxes are not leads
    assert client.get("/api/export/all/counts").json()["prospects"] == 1


def test_a_pitch_mailbox_never_becomes_an_application_address_of_a_listing_of_the_same_company(db, client, fake_site):
    from .helpers import FakeAdapter, raw
    from .test_csv_and_runs import go
    go(db, {"fake": FakeAdapter([raw(1, company_name="Northwind Docs", company_website="https://northwind.example",
                                     description="Python FastAPI LLM AWS Docker work. To apply, email your resume to jobs@northwind.example. " * 3)])})
    client.post("/api/prospects", json={"items": [{"website": "https://northwind.example", "name": "Northwind Docs"}]})
    assert db.query(Company).count() == 1                                              # one company: the listing's and the prospect's
    wb = load_workbook(io.BytesIO(exporter.build_xlsx(exporter.build_combined(db))))
    leads = {r[3].value for r in wb["Leads"].iter_rows(min_row=2)}
    assert leads == {"jobs@northwind.example"}                                          # the site's general mailboxes are not job leads
    head = [c.value for c in wb["Emails"][1]]
    rows = [dict(zip(head, [c.value for c in r])) for r in wb["Emails"].iter_rows(min_row=2)]
    by_email = {r["business_email"]: r for r in rows}
    assert by_email["hello@northwind.example"]["contact_purpose"] == "b2b_mailbox"      # visible, with its own purpose
    assert by_email["jobs@northwind.example"]["contact_purpose"] != "b2b_mailbox"
    assert client.get("/api/overview").json()["emails_found"] == 1


def test_rescoring_and_ease_refresh_leave_prospects_alone(db):
    from app.services import leads, rescoring
    o = P.add_prospect(db, "https://northwind.example", "Northwind Docs", "company")
    P.research_prospect(db, o, site())
    db.refresh(o)
    before = (o.ease_score, list(o.ease_signals), o.fit_score, o.recommended_action)
    rescoring.rescore_all(db); leads.refresh_ease(db); rescoring.refresh_hn_locations(db)
    db.refresh(o)
    assert (o.ease_score, list(o.ease_signals), o.fit_score, o.recommended_action) == before


def test_migration_0004_adds_meta_with_a_default(tmp_path):
    import sqlite3
    from alembic import command
    from alembic.config import Config
    root = __import__("pathlib").Path(__file__).resolve().parents[1]
    url = f"sqlite:///{tmp_path / 'm.db'}"
    cfg = Config(str(root / "alembic.ini")); cfg.set_main_option("script_location", str(root / "migrations"))
    cfg.cmd_opts = type("o", (), {"x": [f"url={url}"]})()
    command.upgrade(cfg, "0003")
    con = sqlite3.connect(tmp_path / "m.db")
    con.execute("insert into opportunities (source_key,source_type,external_id,original_url,canonical_url,fingerprint,content_hash,title,company_name,description,required_skills,matched_skills,missing_requirements,engagement_type,remote_status,location_requirements,currency,rate_type,first_seen_at,last_seen_at,listing_status,fit_score,score_components,fit_explanation,portfolio_evidence,recommended_action,review_status,exclusion_reason,research_status,proposal_status,application_status,ease_score,ease_signals,apply_method) values ('x','api','1','u','u','f','','t','c','d','[]','[]','[]','unknown','unknown','','','','2026-01-01','2026-01-01','open',0,'{}','','[]','','new','','none','none','none',0,'[]','unknown')")
    con.commit(); con.close()
    command.upgrade(cfg, "head")
    con = sqlite3.connect(tmp_path / "m.db")
    assert con.execute("select meta from opportunities").fetchone()[0] == "{}"
    command.downgrade(cfg, "0003")


def test_the_one_file_download_works_with_prospects_alone_but_the_flat_csv_needs_opportunities(client, db):
    assert client.get("/api/export/all").status_code == 404 and "add a prospect" in client.get("/api/export/all").json()["detail"]
    client.post("/api/prospects", json={"items": [{"website": "https://northwind.example", "name": "Northwind Docs", "evidence_quote": "Sells delivery capacity"}], "research": False})
    r = client.get("/api/export/all")
    assert r.status_code == 200 and "X-Counts" in r.headers and "prospects=1" in r.headers["X-Counts"] and "runs=0" in r.headers["X-Counts"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Leads", "Opportunities", "Prospects", "Emails", "Runs"] and wb["Prospects"].max_row == 2
    c = client.get("/api/export/all?format=csv")
    assert c.status_code == 404 and "Excel file also holds your prospects" in c.json()["detail"]


@pytest.mark.parametrize("quote", ['=HYPERLINK("http://evil.example","click")', "+SUM(1,1)", "-2+3", "@cmd", "\t=1+1"])
def test_text_typed_into_a_prospect_never_becomes_a_formula_in_the_downloads(client, db, quote):
    pid = client.post("/api/prospects", json={"items": [{"website": "https://northwind.example", "name": quote[:40], "evidence_quote": quote}], "research": False}).json()["created"][0]["id"]
    assert pid
    r = client.get("/api/export/all")
    wb = load_workbook(io.BytesIO(r.content))
    head = [c.value for c in wb["Prospects"][1]]
    cells = {h: c for h, c in zip(head, wb["Prospects"][2])}
    for key in ("company", "why_added"):
        c = cells[key]
        assert c.data_type == "s"                                              # a text cell: Excel never evaluates it, whatever it starts with
    assert quote.strip() in str(cells["why_added"].value) or quote in str(cells["why_added"].value)    # and the typed text is kept as typed
